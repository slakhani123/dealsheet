#!/usr/bin/env python3
"""Build Contact_List_REL.xlsx from automation/contacts_data.json.

The JSON file is the machine-readable source of truth maintained by the weekly
Outlook sweep; this script renders it into the shareable workbook. Run from the
repo root:

    python3 automation/build_contacts_xlsx.py

Requires openpyxl. The workbook contains COUNTIF formulas on the Summary tab,
so recalculate after building if a tool needs cached values.
"""

import json
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "automation", "contacts_data.json")
DEFAULT_OUT = os.path.join(ROOT, "Contact_List_REL.xlsx")

CATEGORY_ORDER = [
    "Broker",
    "Lender / Bank",
    "Borrower / Sponsor",
    "Investor / Family Office",
    "Solicitor / Legal",
    "Valuer / Property Professional",
    "Accountant / Tax",
    "Service Provider",
    "Internal",
    "Personal / Other",
]

NAVY = "1F3864"
LIGHT = "D9E2F3"
FONT = "Arial"

CAT_FILL = {
    "Broker": "FCE4D6",
    "Lender / Bank": "DDEBF7",
    "Borrower / Sponsor": "E2EFDA",
    "Investor / Family Office": "FFF2CC",
    "Solicitor / Legal": "EDEDED",
    "Valuer / Property Professional": "D9E1F2",
    "Accountant / Tax": "FBE5EB",
    "Service Provider": "F2F2F2",
    "Internal": "C6E0B4",
    "Personal / Other": "FFFFFF",
}


def main():
    with open(DATA) as f:
        payload = json.load(f)
    cfg = payload.get("config") or {}
    org = cfg.get("org_name") or ""
    owner = cfg.get("owner_name") or ""
    owner_first = cfg.get("owner_first_name") or owner.split(" ")[0] or "the list owner"
    owner_email = cfg.get("owner_email") or ""
    title_txt = f"{org} — Contact List" if org else "Contact List"
    # Filename is configurable so several mailboxes can live side by side.
    out_path = os.path.join(ROOT, cfg.get("output_basename") or
                            os.path.basename(DEFAULT_OUT))
    contacts = [c for c in payload["contacts"] if c.get("keep", True)]
    refreshed = payload.get("last_sweep_utc", "")[:10] or date.today().isoformat()

    order = {c: i for i, c in enumerate(CATEGORY_ORDER)}
    contacts.sort(key=lambda c: (
        order.get(c.get("category", "Personal / Other"), 99),
        (c.get("company") or "zzz").lower(),
        (c.get("name") or c["email"]).lower(),
    ))

    wb = Workbook()
    # openpyxl writes formulas without cached values; force the reader to
    # evaluate the Summary tab's COUNTIFs the moment the file opens.
    wb.calculation.fullCalcOnLoad = True

    # ---- Contact List ----
    ws = wb.active
    ws.title = "Contact List"
    headers = ["Name", "Category", "Company", "Job Title", "Email", "Phone",
               "First Seen", "Last Seen", "Emails In", "Emails Out",
               "How Found", "Notes"]
    widths = [24, 24, 30, 26, 34, 24, 12, 12, 10, 11, 19, 58]

    ws.merge_cells("A1:L1")
    t = ws["A1"]
    t.value = title_txt
    t.font = Font(name=FONT, size=16, bold=True, color=NAVY)
    ws.merge_cells("A2:L2")
    s = ws["A2"]
    s.value = (f"Built automatically from Outlook"
               + (f" ({owner_email})" if owner_email else "")
               + f" — last refreshed {refreshed}. Refreshed weekly; "
               f"do not hand-edit (see README tab).")
    s.font = Font(name=FONT, size=9, italic=True, color="666666")

    hrow = 4
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=hrow, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(i)].width = w

    r = hrow
    for c in contacts:
        r += 1
        cat = c.get("category", "Personal / Other")
        vals = [
            c.get("name") or c["email"],
            cat,
            c.get("company", ""),
            c.get("job_title", ""),
            c["email"],
            ", ".join(c.get("phones", [])),
            c.get("first_seen", ""),
            c.get("last_seen", ""),
            c.get("count_from", 0),
            c.get("count_to", 0),
            ("Direct correspondence"
             if (c.get("count_from", 0) or c.get("count_to", 0))
             else "Named in a thread"),
            c.get("note", ""),
        ]
        for i, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.border = border
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(i == 12),
                                       horizontal="center" if i in (7, 8, 9, 10) else "left")
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=CAT_FILL.get(cat, "FFFFFF"))
        if vals[10] == "Named in a thread":
            ws.cell(row=r, column=11).font = Font(name=FONT, size=10, italic=True,
                                                  color="808080")

    ws.auto_filter.ref = f"A{hrow}:L{r}"
    ws.freeze_panes = f"A{hrow + 1}"

    # ---- Summary ----
    sm = wb.create_sheet("Summary")
    sm.merge_cells("A1:C1")
    sm["A1"] = "Contacts by category"
    sm["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    sm["A3"] = "Category"
    sm["B3"] = "Contacts"
    for col in ("A3", "B3"):
        sm[col].font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        sm[col].fill = PatternFill("solid", fgColor=NAVY)
        sm[col].border = border
    sm.column_dimensions["A"].width = 32
    sm.column_dimensions["B"].width = 12
    row = 3
    for cat in CATEGORY_ORDER:
        row += 1
        a = sm.cell(row=row, column=1, value=cat)
        b = sm.cell(row=row, column=2,
                    value=f"=COUNTIF('Contact List'!$B$5:$B${r},A{row})")
        a.font = Font(name=FONT, size=10)
        b.font = Font(name=FONT, size=10)
        a.fill = PatternFill("solid", fgColor=CAT_FILL.get(cat, "FFFFFF"))
        a.border = border
        b.border = border
        b.alignment = Alignment(horizontal="center")
    row += 1
    ta = sm.cell(row=row, column=1, value="Total")
    tb = sm.cell(row=row, column=2, value=f"=SUM(B4:B{row - 1})")
    ta.font = Font(name=FONT, size=10, bold=True)
    tb.font = Font(name=FONT, size=10, bold=True)
    ta.border = border
    tb.border = border
    tb.alignment = Alignment(horizontal="center")

    # ---- README ----
    rd = wb.create_sheet("README")
    rd.column_dimensions["A"].width = 110
    lines = [
        (title_txt, 14, True),
        ("", 10, False),
        ("What this is", 11, True),
        (f"Every external (and internal) contact from {owner or 'the'} Outlook mailbox, "
         "categorised as Broker, Lender / Bank, Borrower / Sponsor, Solicitor / Legal, "
         "Valuer, Investor, Accountant, Service Provider, Internal or Other.", 10, False),
        ("", 10, False),
        ("How it is produced", 11, True),
        ("A weekly automated sweep reads the mailbox (Inbox, Sent Items and Archive), "
         "extracts every correspondent plus signature details (company, job title, phone), "
         "categorises them from deal context, and rebuilds this workbook. The master data "
         "lives in the dealsheet git repository (automation/contacts_data.json); this file "
         "is regenerated and re-uploaded to SharePoint automatically on every refresh.", 10, False),
        ("", 10, False),
        ("Rules of use", 11, True),
        ("1. Do NOT hand-edit this file in SharePoint — the weekly refresh overwrites it. "
         f"Corrections (wrong category, name, company etc.) should go to {owner_first}, who will "
         "apply them to the master data so they stick.", 10, False),
        ("2. 'Emails In / Out' are total messages received from / sent to that contact "
         "since Jan 2025 — a rough measure of relationship activity.", 10, False),
        ("3. 'How Found' says where the address came from. 'Direct correspondence' means "
         f"the contact has emailed {owner_first} or been emailed by them. 'Named in a thread' means "
         "the address was recovered from a forwarded message header or a quoted signature "
         "— the contact is real, but nobody at REL has emailed them from this mailbox.", 10, False),
        ("4. Automated senders (newsletters, notifications, no-reply addresses) are "
         "excluded automatically.", 10, False),
        ("5. Phone numbers and job titles are mined from email signatures where available; "
         "blanks mean no signature was found, not that the contact has none.", 10, False),
    ]
    for i, (txt, size, bold) in enumerate(lines, start=1):
        c = rd.cell(row=i, column=1, value=txt)
        c.font = Font(name=FONT, size=size, bold=bold, color=NAVY if bold else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(out_path)
    print(f"Wrote {out_path}: {len(contacts)} contacts")


if __name__ == "__main__":
    sys.exit(main())
